from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List

from docx import Document
from docx.shared import Pt, Inches


DATE_FMT = "%d %B %Y"


def parse_date(value: str | None) -> datetime | None:
    if not value:
        return None
    value = value.strip()
    for fmt in (DATE_FMT, "%d %b %Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def format_date(value: str | None) -> str:
    dt = parse_date(value)
    return dt.strftime(DATE_FMT) if dt else (value or "")


def calc_age(dob: str | None, on_date: str | None) -> str:
    dob_dt = parse_date(dob)
    ref_dt = parse_date(on_date)
    if not dob_dt or not ref_dt:
        return ""
    years = ref_dt.year - dob_dt.year - ((ref_dt.month, ref_dt.day) < (dob_dt.month, dob_dt.day))
    return f"{years} years"


def first_name(full_name: str | None) -> str:
    return (full_name or "").strip().split()[0] if (full_name or "").strip() else ""


def sanitize_filename(value: str) -> str:
    cleaned = value.replace("/", "_").replace("\\", "_").replace(" ", "_")
    return "".join(ch for ch in cleaned if ch.isalnum() or ch in "_-.") or "case"


def set_cell_text(cell, text: str):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text or "")
    run.font.size = Pt(10)




def clear_cell_shading(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('}shd'):
            tcPr.remove(child)


def clear_doc_highlighting(doc: Document):
    for paragraph in iter_all_paragraphs(doc):
        for run in paragraph.runs:
            try:
                run.font.highlight_color = None
            except Exception:
                pass
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                clear_cell_shading(cell)


def ensure_table_columns(table, total_columns: int):
    while len(table.columns) < total_columns:
        table.add_column(Inches(1.2))


def normalize_row_lengths_variable(rows: List[List[str]], width: int | None = None) -> List[List[str]]:
    rows = rows or []
    if width is None:
        width = max([len(r) for r in rows], default=0)
    out = []
    for row in rows:
        vals = list(row[:width]) + [''] * max(0, width - len(row))
        if any((v or '').strip() for v in vals):
            out.append(vals)
    return out

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
    elif level == 2:
        run.font.size = Pt(12)
    else:
        run.font.size = Pt(11)
    return p


def add_body(doc, text):
    for part in (text or "").split("\n"):
        doc.add_paragraph(part.strip())


def format_edu_rows(rows):
    if not rows:
        return "Not provided."
    out = []
    for row in rows:
        if any((x or "").strip() for x in row):
            out.append(" | ".join([x for x in row if x]))
    return "\n".join(out) if out else "Not provided."


def list_from_multiline(value: str | None) -> List[str]:
    if not value:
        return []
    lines = []
    for raw in value.splitlines():
        cleaned = raw.strip().lstrip("•-").strip()
        if cleaned:
            lines.append(cleaned)
    return lines


def normalize_row_lengths(rows: List[List[str]], width: int) -> List[List[str]]:
    out = []
    for row in rows or []:
        vals = list(row[:width]) + [""] * max(0, width - len(row))
        if any((v or "").strip() for v in vals):
            out.append(vals[:width])
    return out


def parse_delimited_rows(value: str | None, expected_cols: int) -> List[List[str]]:
    if not value:
        return []
    rows: List[List[str]] = []
    for raw in value.splitlines():
        line = raw.strip()
        if not line:
            continue
        parts = [part.strip() for part in line.split('|')]
        if len(parts) == 1:
            parts = [part.strip() for part in line.split(';')]
        if len(parts) == 1:
            parts = [part.strip() for part in line.split(',')]
        parts = parts[:expected_cols] + [""] * max(0, expected_cols - len(parts))
        if any(parts):
            rows.append(parts[:expected_cols])
    return rows


def narrative_summary(data):
    claimant = data.get("claimant_full_name") or data.get("name", "The claimant")
    gender = (data.get("gender") or "").lower()
    pronoun = "The claimant"
    if gender == "male":
        pronoun = "He"
    elif gender == "female":
        pronoun = "She"

    parts = [
        f"{claimant} was assessed on {data.get('assessment_date', '')} at {data.get('venue_evaluation', '')}. The consultation was conducted by {data.get('assessor', '')}.",
    ]
    if data.get("family_history"):
        parts.append(f"Family history summary: {data.get('family_history')}")
    if data.get("accident_detail"):
        parts.append(f"Accident summary: {data.get('accident_detail')}")
    complaints = []
    if data.get("education_difficulties_client"):
        complaints.append(f"educational difficulties including {data.get('education_difficulties_client')}")
    if data.get("physical_difficulties"):
        complaints.append(f"physical difficulties including {data.get('physical_difficulties')}")
    if data.get("cognitive_difficulties"):
        complaints.append(f"cognitive difficulties including {data.get('cognitive_difficulties')}")
    if data.get("psychological_difficulties"):
        complaints.append(f"psychological difficulties including {data.get('psychological_difficulties')}")
    if complaints:
        parts.append(f"{pronoun} reported " + "; ".join(complaints) + ".")
    if data.get("people_in_house"):
        parts.append(f"Current living arrangement: {data.get('people_in_house')}.")
    return "\n\n".join(parts)


# -------- Intake report generation --------

def fill_table_0(table, d):
    values = [
        d.get("name"), d.get("id_passport"), d.get("dob"), d.get("nationality"),
        d.get("gender"), d.get("place_of_birth"), d.get("home_language"),
        d.get("other_languages"), d.get("highest_education"),
        d.get("pre_accident_education"), d.get("current_education"),
        d.get("physical_address"), d.get("client_contact"), d.get("email"),
        d.get("accompanying_person"), d.get("date_accident"),
        d.get("date_evaluation"), d.get("venue_evaluation"),
    ]
    for i, value in enumerate(values):
        set_cell_text(table.cell(i, 1), value or "")


def fill_table_1(table, d):
    set_cell_text(table.cell(1, 0), d.get("family_history", ""))


def fill_table_2(table, d):
    primary_rows = normalize_row_lengths_variable(d.get("primary_education") or [], 4)
    high_rows = normalize_row_lengths_variable(d.get("highschool_education") or [], 4)
    if not primary_rows and not high_rows:
        all_rows = normalize_row_lengths_variable(d.get("education_history_all") or [], 4)
        if all_rows:
            primary_rows = all_rows[:4]
            high_rows = all_rows[4:]

    # Detect the actual row positions from the uploaded intake template instead of assuming fixed indexes.
    primary_start = 3
    high_label_row = None
    high_start = None
    for idx, row in enumerate(table.rows):
        row_text = ' '.join((cell.text or '').strip().lower() for cell in row.cells)
        if 'primary school' in row_text:
            primary_start = idx + 2 if idx + 2 < len(table.rows) else idx + 1
        if 'high school' in row_text:
            high_label_row = idx
            high_start = idx + 2 if idx + 2 < len(table.rows) else idx + 1

    if high_start is None:
        high_start = max(primary_start + max(1, len(primary_rows)) + 2, 6)

    required_rows = max(len(table.rows), primary_start + max(1, len(primary_rows)), high_start + max(1, len(high_rows)))
    ensure_table_rows(table, required_rows)

    # Clear existing editable rows in the primary section until the next section label.
    primary_clear_end = (high_label_row or high_start) - 1 if high_label_row is not None else high_start - 1
    for ridx in range(primary_start, max(primary_start, primary_clear_end)):
        if ridx < len(table.rows):
            for c in range(min(4, len(table.columns))):
                set_cell_text(table.cell(ridx, c), '')
    for ridx, row in enumerate(primary_rows, start=primary_start):
        for c in range(4):
            set_cell_text(table.cell(ridx, c), row[c] if c < len(row) else '')

    # Clear existing editable rows in the high school section and then populate.
    for ridx in range(high_start, len(table.rows)):
        for c in range(min(4, len(table.columns))):
            set_cell_text(table.cell(ridx, c), '')
    for ridx, row in enumerate(high_rows, start=high_start):
        if ridx >= len(table.rows):
            ensure_table_rows(table, ridx + 1)
        for c in range(4):
            set_cell_text(table.cell(ridx, c), row[c] if c < len(row) else '')

def fill_table_3(table, d):
    vals = [d.get("grade_repeated_before"), d.get("grade_at_accident"), d.get("grade_repeated_after"), d.get("highest_grade_passed")]
    for i, v in enumerate(vals):
        set_cell_text(table.cell(i, 1), v or "")


def fill_table_4(table, d):
    set_cell_text(table.cell(1, 0), "Pre-Accident:\n" + (d.get("future_plans_pre") or ""))
    set_cell_text(table.cell(2, 0), "Post-Accident:\n" + (d.get("future_plans_post") or ""))


def fill_table_5(table, d):
    mapping = {
        (1, 1): d.get("accident_detail"),
        (2, 1): d.get("transport_hospital"),
        (3, 1): d.get("injuries_sustained"),
        (4, 1): d.get("treatment_administered"),
        (5, 1): d.get("hospital_duration"),
        (6, 1): "Client (provided they can speak for themselves):\n" + (d.get("education_difficulties_client") or "") + "\n\nGuardian:\n" + (d.get("education_difficulties_guardian") or ""),
        (7, 1): "Physical difficulties:\n" + (d.get("physical_difficulties") or ""),
        (8, 1): "Cognitive difficulties:\n" + (d.get("cognitive_difficulties") or ""),
        (9, 1): "Psychological difficulties:\n" + (d.get("psychological_difficulties") or ""),
        (10, 1): d.get("behaviour_observation"),
        (13, 1): d.get("schooling_at_accident"),
        (14, 1): d.get("off_school_duration"),
        (15, 1): d.get("returned_to_school"),
        (16, 1): d.get("disability_grant"),
        (17, 1): d.get("financial_dependence"),
        (18, 1): d.get("documentation"),
    }
    for (r, c), text in mapping.items():
        set_cell_text(table.cell(r, c), text or "")


def fill_table_6(table, d):
    vals = [
        d.get("housing_type"), d.get("housing_ownership"), d.get("bedrooms"),
        d.get("dining_room"), d.get("kitchen"), d.get("bathroom"),
        d.get("toilet"), d.get("outside_rooms"), d.get("people_in_house"),
        d.get("water"), d.get("electricity"), d.get("amenities_distance"),
        d.get("transport_mode")
    ]
    for idx, val in enumerate(vals, start=1):
        set_cell_text(table.cell(idx, 1), val or "")


def fill_table_7(table, d):
    set_cell_text(table.cell(1, 0), d.get("collateral_info", ""))


def fill_table_8(table, d):
    vals = [
        d.get("claimant_full_name"), d.get("reference"), d.get("assessment_date"),
        d.get("assessor"), d.get("followup_assessor"), d.get("attorney_name"),
        d.get("attorney_reference")
    ]
    for idx, val in enumerate(vals, start=1):
        set_cell_text(table.cell(idx, 1), val or "")


def generate_report_docx(template_path, output_path, data):
    template_path = Path(template_path)
    output_path = Path(output_path)
    doc = Document(str(template_path))

    fillers = [fill_table_0, fill_table_1, fill_table_2, fill_table_3, fill_table_4, fill_table_5, fill_table_6, fill_table_7, fill_table_8]
    for table, filler in zip(doc.tables, fillers):
        filler(table, data)

    doc.add_page_break()
    add_heading(doc, "Narrative Intake Summary", level=1)
    add_body(doc, narrative_summary(data))
    add_heading(doc, "Educational History Summary", level=2)
    add_body(doc, "Primary School:\n" + format_edu_rows(data.get("primary_education")))
    add_body(doc, "High School:\n" + format_edu_rows(data.get("highschool_education")))

    clear_doc_highlighting(doc)
    doc.save(str(output_path))


# -------- Industrial Psychology report generation --------

def iter_all_paragraphs(doc: Document):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_markers_in_paragraph(paragraph, replacements: Dict[str, str]):
    combined = "".join(run.text for run in paragraph.runs)
    if not combined or "[[" not in combined:
        return
    replaced = combined
    for marker, value in replacements.items():
        replaced = replaced.replace(f"[[{marker}]]", value or "")
    if replaced == combined:
        return
    for idx in range(len(paragraph.runs) - 1, -1, -1):
        paragraph.runs[idx]._element.getparent().remove(paragraph.runs[idx]._element)
    paragraph.add_run(replaced)


def ensure_table_rows(table, total_rows: int):
    while len(table.rows) < total_rows:
        new_row = table.add_row()
        if len(table.rows) > 1:
            src_tr = table.rows[-2]._tr
            new_tr = new_row._tr
            for src_tc, new_tc in zip(src_tr.tc_lst, new_tr.tc_lst):
                tc_pr = src_tc.tcPr
                if tc_pr is not None:
                    new_tc.insert(0, deepcopy(tc_pr))


def fill_client_info_table(table, data: dict):
    values = {
        1: data.get("name", ""),
        2: data.get("id_passport", ""),
        3: format_date(data.get("dob")),
        4: data.get("nationality", ""),
        5: data.get("place_of_birth", ""),
        6: data.get("current_age", ""),
        7: data.get("gender", ""),
        8: data.get("home_language", ""),
        9: data.get("highest_education", ""),
        10: format_date(data.get("date_accident")),
        11: data.get("age_at_accident", ""),
        12: data.get("grade_at_accident", ""),
        13: data.get("current_grade", data.get("current_education", "")),
        14: data.get("type_of_accident", ""),
        15: format_date(data.get("assessment_date")),
        16: format_date(data.get("report_date")),
        17: data.get("attorney_name", ""),
        18: data.get("attorney_reference", ""),
        19: data.get("our_reference", ""),
    }
    for row_idx, value in values.items():
        set_cell_text(table.cell(row_idx, 1), value)


def fill_dynamic_table(table, rows: List[List[str]], start_row: int = 1, preserve_header: bool = True, headers: List[str] | None = None):
    rows = normalize_row_lengths_variable(rows)
    target_cols = max(len(table.columns), len(headers or []), max((len(r) for r in rows), default=len(table.columns)))
    ensure_table_columns(table, target_cols)
    required_rows = start_row + max(1, len(rows))
    ensure_table_rows(table, required_rows)
    if preserve_header and headers:
        for col_idx, header in enumerate(headers):
            set_cell_text(table.cell(start_row - 1, col_idx), header or "")
    for row_idx in range(start_row, len(table.rows)):
        current_values = rows[row_idx - start_row] if row_idx - start_row < len(rows) else ["" for _ in range(target_cols)]
        for col_idx in range(target_cols):
            value = current_values[col_idx] if col_idx < len(current_values) else ""
            set_cell_text(table.cell(row_idx, col_idx), value)


def first_non_empty(*values: str) -> str:
    for value in values:
        if value and str(value).strip():
            return str(value).strip()
    return ""


def ensure_sentence(value: str, fallback: str = "") -> str:
    text = (value or fallback or "").strip()
    if not text:
        return ""
    if text[-1] not in ".!?":
        text += "."
    return text


def title_case_relation(value: str | None) -> str:
    text = (value or "").strip()
    return text.title() if text else ""


def summarize_people_in_house(value: str | None) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    return text


def build_residential_paragraph(d: dict) -> str:
    claimant = d.get("claimant_short_name") or "The claimant"
    residents = summarize_people_in_house(d.get("people_in_house"))
    house_type = (d.get("housing_type") or "").strip()
    bedrooms = (d.get("bedrooms") or "").strip()
    dining = (d.get("dining_room") or "").strip()
    kitchen = (d.get("kitchen") or "").strip()
    bathroom = (d.get("bathroom") or "").strip()
    toilet = (d.get("toilet") or "").strip()
    water = (d.get("water") or "").strip()
    electricity = (d.get("electricity") or "").strip()
    transport = (d.get("transport_mode") or "").strip()
    amenities = (d.get("amenities_distance") or "").strip()

    bits = []
    if residents:
        bits.append(f"{claimant} resides with {residents}")
    elif d.get("guardian_name"):
        bits.append(f"{claimant} resides with {d.get('guardian_name')}")
    if house_type:
        house = house_type
        if bedrooms:
            house = f"a {bedrooms}-bedroomed {house_type}"
        bits.append(f"They live in {house}")
    rooms = []
    if dining:
        rooms.append(dining)
    if kitchen:
        rooms.append(kitchen)
    if bathroom:
        rooms.append(bathroom)
    if rooms:
        bits.append("The dwelling includes " + ", ".join(rooms))
    if toilet:
        bits.append(f"They use {toilet}")
    if water:
        bits.append(f"Water access is described as {water}")
    if electricity:
        bits.append(f"Electricity access is described as {electricity}")
    if amenities:
        bits.append(f"Amenities are reported as {amenities}")
    if transport:
        bits.append(f"Transport is mainly {transport}")

    sentences = [ensure_sentence(b).rstrip() for b in bits if b.strip()]
    paragraph = " ".join(sentences)
    return ensure_sentence(paragraph, f"{claimant}'s current residential circumstances were recorded during intake")


def build_medical_history_paragraph(d: dict) -> str:
    guardian = d.get("guardian_reference_name") or d.get("guardian_name") or "The accompanying informant"
    claimant = d.get("claimant_short_name") or "the claimant"
    other_accidents = first_non_empty(d.get("other_accidents_count"), "no")
    developmental = first_non_empty(d.get("developmental_milestones"), "were reportedly achieved within the expected developmental timeframes")
    prior_history = first_non_empty(d.get("medical_history_prior"), "no significant pre-existing medical history was reported")
    text = (
        f"{guardian} reported that the accident under discussion was {claimant}'s "
        f"{first_non_empty(d.get('accident_ordinal'), 'first')} accident and that "
        f"{('she' if d.get('claimant_pronoun_object') == 'her' else 'he')} had been involved in {other_accidents} other accidents thereafter. "
        f"Prior to the accident, {prior_history}. Developmental milestones {developmental}."
    )
    return ensure_sentence(text)


def build_education_note(d: dict) -> str:
    rows = d.get("education_history_all") or []
    if rows:
        institutions = []
        for row in rows[:3]:
            if row and row[0]:
                institutions.append(row[0])
        institution_text = ", ".join(dict.fromkeys(institutions))
        if institution_text:
            return ensure_sentence(
                f"The above school progression is based on the intake information and the school history captured for this matter, including records from {institution_text}"
            )
    return ensure_sentence(first_non_empty(
        d.get("education_nb_text"),
        "School progression is based on the information captured during intake and any academic records made available"
    ))


def build_collateral_paragraph(d: dict) -> str:
    claimant = d.get("claimant_short_name") or "the claimant"
    pre = (d.get("pre_accident_educator") or "").strip()
    current = (d.get("current_educator") or "").strip()
    principal = (d.get("principal_info") or "").strip()
    contacts = (d.get("collateral_info") or "").strip()
    bits = [
        "Obtaining collateral information is crucial for understanding the full impact of the accident on the claimant's educational attainment."
    ]
    source_bits = []
    if pre:
        source_bits.append(f"pre-accident educator: {pre}")
    if current:
        source_bits.append(f"current educator: {current}")
    if principal:
        source_bits.append(f"principal information: {principal}")
    if contacts:
        source_bits.append(f"other verification contacts: {contacts}")
    if source_bits:
        bits.append(f"For {claimant}, the available collateral sources include " + "; ".join(source_bits) + ".")
    else:
        bits.append(f"No direct collateral contacts were captured at the time of drafting for {claimant}.")
    return " ".join(bits)


def build_accident_paragraph(d: dict) -> str:
    guardian = d.get("guardian_reference_name") or d.get("guardian_name") or "The informant"
    date = format_date(d.get("date_accident"))
    detail = first_non_empty(d.get("accident_detail"))
    sentence = detail.strip() if detail else f"On {date}, the claimant was involved in an accident."
    transport = first_non_empty(d.get("transport_hospital"))
    if transport:
        sentence += " " + ensure_sentence(transport)
    return ensure_sentence(f"{guardian} reported that {sentence}")


def build_behaviour_paragraph(d: dict) -> str:
    claimant = d.get("claimant_short_name") or "The claimant"
    guardian = d.get("guardian_name") or d.get("guardian_reference_name") or ""
    arrival = (d.get("arrival_time") or "").strip()
    observed = (d.get("behaviour_observation") or "").strip()
    bits = []
    if arrival:
        bits.append(f"{claimant} arrived {arrival} for the scheduled assessment")
    else:
        bits.append(f"{claimant} attended the scheduled assessment")
    if guardian:
        bits[-1] += f", accompanied by {guardian}"
    if observed:
        bits.append(observed)
    return ensure_sentence(". ".join([b.strip().rstrip('.') for b in bits]))


def build_recommendation_paragraphs(d: dict) -> tuple[str, str]:
    complaint_text = first_non_empty(
        d.get("industrial_psych_summary"),
        d.get("post_morbid_potential"),
        d.get("psychological_difficulties"),
        d.get("cognitive_difficulties"),
        d.get("physical_difficulties"),
    )
    if complaint_text:
        lowered = complaint_text[0].lower() + complaint_text[1:] if complaint_text else complaint_text
        p1 = ensure_sentence(f"Following the accident, the minor continues to experience {lowered}")
    else:
        p1 = "Following the accident, the available information suggests persisting functional and educational difficulties."
    p2 = "Therefore, provision should be made for the recommended interventions and treatments proposed by the experts."
    return p1, p2


def set_paragraph_text(paragraph, text: str):
    for idx in range(len(paragraph.runs) - 1, -1, -1):
        paragraph.runs[idx]._element.getparent().remove(paragraph.runs[idx]._element)
    paragraph.text = ""
    paragraph.add_run(text or "")


def overwrite_paragraph_after(doc: Document, heading_text: str, new_text: str):
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == heading_text:
            for j in range(idx + 1, len(paragraphs)):
                if paragraphs[j].text.strip():
                    set_paragraph_text(paragraphs[j], new_text)
                    return


def overwrite_document_received_lines(doc: Document, lines: List[str]):
    paragraphs = doc.paragraphs
    start = end = None
    for idx, paragraph in enumerate(paragraphs):
        txt = paragraph.text.strip()
        if txt == "The following supporting documentation was received:":
            start = idx + 1
        elif start is not None and txt == "BACKGROUND INFORMATION":
            end = idx
            break
    if start is None or end is None:
        return
    target_count = end - start
    items = lines[:target_count] + [""] * max(0, target_count - len(lines))
    for offset, value in enumerate(items):
        set_paragraph_text(paragraphs[start + offset], value)


def build_documents_received(d: dict) -> List[str]:
    return [
        f"Instruction Letter from {first_non_empty(d.get('attorney_name'), 'the instructing attorneys')}",
        f"Medical Records from {first_non_empty(d.get('medical_records_from'), d.get('transport_hospital'), 'the available hospitals and medical facilities')}",
        f"RAF 01 Form (Medical Practitioner) completed by {first_non_empty(d.get('raf1_completed_by'), 'the available medical practitioner')}",
        f"RAF 04 Form (Orthopaedic Surgeon) completed by {first_non_empty(d.get('raf4_completed_by'), 'the orthopaedic surgeon')}",
        f"Orthopaedic Surgeon Report completed by {first_non_empty(d.get('orthopaedic_report_by'), 'the orthopaedic surgeon')}",
        f"Diagnostic Radiologist Report completed by {first_non_empty(d.get('radiologist_report_by'), 'the diagnostic radiologist')}",
        f"Educational Psychologist Report completed by {first_non_empty(d.get('educational_psychologist_by'), 'the educational psychologist')}",
        f"Occupational Therapist Report completed by {first_non_empty(d.get('occupational_therapist_by'), 'the occupational therapist')}",
        f"Copies of {first_non_empty(d.get('claimant_short_name'), 'the claimant')}'s School Reports",
        f"Copy of {first_non_empty(d.get('claimant_short_name'), 'the claimant')}'s Birth Certificate",
        f"Copy of {first_non_empty(d.get('guardian_reference_name'), d.get('guardian_name'), 'the guardian')}'s Identity Document",
    ]


def derive_ip_defaults(data: dict) -> dict:
    derived = dict(data)
    claimant_name = first_non_empty(derived.get("claimant_full_name"), derived.get("name"))
    derived.setdefault("claimant_full_name", claimant_name)
    derived.setdefault("name", claimant_name)
    derived.setdefault("claimant_short_name", first_name(claimant_name))
    derived.setdefault("assessment_date", first_non_empty(derived.get("assessment_date"), derived.get("date_evaluation")))
    derived.setdefault("report_date", first_non_empty(derived.get("report_date"), derived.get("assessment_date"), derived.get("date_evaluation")))
    derived.setdefault("current_grade", first_non_empty(derived.get("current_grade"), derived.get("current_education")))
    derived.setdefault("type_of_accident", first_non_empty(derived.get("type_of_accident"), "Pedestrian Vehicle Accident"))

    guardian_raw = first_non_empty(derived.get("accompanying_person"), derived.get("guardian_name"))
    guardian_name = guardian_raw.split("(")[0].strip() if guardian_raw else ""
    derived.setdefault("guardian_name", guardian_name)
    guardian_relation = first_non_empty(derived.get("guardian_relation"), derived.get("relationship_to_client"), "guardian")
    derived.setdefault("guardian_relation", guardian_relation)
    guardian_ref = first_non_empty(derived.get("guardian_reference_name"), guardian_name, title_case_relation(guardian_relation))
    derived.setdefault("guardian_reference_name", guardian_ref)
    derived.setdefault("instruction_from", first_non_empty(derived.get("attorney_name"), "the instructing attorneys"))

    gender = (derived.get("gender") or "").strip().lower()
    possessive = "his"
    obj = "him"
    if gender == "female":
        possessive = "her"
        obj = "her"
    elif gender in ("male", "m"):
        possessive = "his"
        obj = "him"
    derived.setdefault("claimant_pronoun_possessive", possessive)
    derived.setdefault("claimant_pronoun_object", obj)

    derived.setdefault("current_age", calc_age(derived.get("dob"), derived.get("report_date") or derived.get("assessment_date")))
    derived.setdefault("age_at_accident", calc_age(derived.get("dob"), derived.get("date_accident")))

    family_headers = derived.get("family_members_headers") or ["Name and Surname", "Relationship", "Age", "Education", "Occupation"]
    family_rows = normalize_row_lengths_variable(derived.get("family_members") or [], len(family_headers))
    if not family_rows:
        family_rows = parse_delimited_rows(derived.get("family_members_text"), len(family_headers))
    if not family_rows and guardian_name:
        family_rows = [[guardian_name, title_case_relation(guardian_relation), "", "", ""]]
    derived["family_members_headers"] = family_headers
    derived["family_members"] = family_rows

    edu_headers = derived.get("education_history_all_headers") or ["Education Institution", "Year", "Qualification", "Comment"]
    edu_rows = normalize_row_lengths_variable(derived.get("education_history_all") or [], len(edu_headers))
    if not edu_rows:
        edu_rows = parse_delimited_rows(derived.get("education_history_text"), 4)
    if not edu_rows:
        for row in (derived.get("primary_education") or []) + (derived.get("highschool_education") or []):
            vals = list(row[:4]) + [""] * max(0, 4 - len(row))
            if any((v or "").strip() for v in vals):
                edu_rows.append([vals[0], vals[2], vals[1], vals[3]])
    derived["education_history_all_headers"] = edu_headers
    derived["education_history_all"] = edu_rows

    injuries = list_from_multiline(derived.get("injuries_sustained"))
    if injuries:
        derived.setdefault("injury_1", injuries[0])

    treatments = list_from_multiline(derived.get("treatment_administered"))
    for idx, value in enumerate(treatments[:5], start=1):
        derived.setdefault(f"treatment_{idx}", ensure_sentence(value))

    complaints = []
    for key in ("physical_difficulties", "psychological_difficulties", "cognitive_difficulties"):
        complaints.extend(list_from_multiline(derived.get(key)))
    for idx, value in enumerate(complaints[:7], start=1):
        derived.setdefault(f"complaint_{idx}", ensure_sentence(value))

    recs = list_from_multiline(first_non_empty(derived.get("recommendations"), derived.get("recommendation_1")))
    if not recs:
        recs = [first_non_empty(derived.get("recommendation_1")), first_non_empty(derived.get("recommendation_2"))]
    recs = [ensure_sentence(r) for r in recs if r]
    if recs:
        derived["recommendation_1"] = recs[0]
        if len(recs) > 1:
            derived["recommendation_2"] = recs[1]

    derived.setdefault("pre_morbid_potential", first_non_empty(derived.get("pre_morbid_potential"), derived.get("pre_morbid"), derived.get("occupation_pre")))
    derived.setdefault("post_morbid_potential", first_non_empty(derived.get("post_morbid_potential"), derived.get("post_morbid"), derived.get("occupation_post")))
    derived.setdefault("loss_of_earnings_text", first_non_empty(derived.get("loss_of_earnings_text"), derived.get("loss_of_earnings_summary")))
    derived.setdefault("education_nb_text", build_education_note(derived))
    derived.setdefault("parents_description", first_non_empty(derived.get("parents_description"), "recorded in the family history table"))
    derived.setdefault("collateral_has_contacts", "were" if derived.get("collateral_info") else "were not")
    return derived


def build_marker_map(data: dict) -> Dict[str, str]:
    d = derive_ip_defaults(data)
    claimant = d.get("claimant_short_name") or first_name(d.get("claimant_full_name")) or "the minor"
    guardian = d.get("guardian_name") or d.get("guardian_reference_name") or "the guardian"
    guardian_ref = d.get("guardian_reference_name") or guardian
    subject_phrase = d.get("guardian_relation_phrase") or f"{d.get('claimant_pronoun_possessive', 'her')} {d.get('guardian_relation', 'guardian')}"
    family_desc = d.get("parents_description") or "noted in the intake information"

    marker_map = {
        "P016_01": claimant,
        "P038_01": d.get("instruction_from", ""),
        "P038_02": guardian,
        "P038_03": guardian_ref,
        "P038_04": subject_phrase,
        "P038_05": claimant,
        "P038_06": claimant,
        "P038_07": claimant,
        "P038_08": claimant,
        "P038_09": d.get("type_of_accident", ""),
        "P038_10": parse_date(d.get("date_accident", "")).strftime("%d") if parse_date(d.get("date_accident", "")) else "",
        "P038_11": parse_date(d.get("date_accident", "")).strftime("%B %Y") if parse_date(d.get("date_accident", "")) else d.get("date_accident", ""),
        "P041_01": claimant,
        "P044_01": format_date(d.get("assessment_date", "")),
        "P047_01": claimant,
        "P049_01": guardian_ref,
        "P049_02": claimant,
        "P052_01": claimant,
        "P052_02": guardian_ref,
        "P052_03": claimant,
        "P056_01": d.get("attorney_name", ""),
        "P057_01": first_non_empty(d.get("medical_records_from"), d.get("transport_hospital")),
        "P058_01": first_non_empty(d.get("raf1_completed_by"), "the available medical practitioner"),
        "P059_01": first_non_empty(d.get("raf4_completed_by"), "the orthopaedic surgeon"),
        "P060_01": first_non_empty(d.get("orthopaedic_report_by"), "the orthopaedic surgeon"),
        "P061_01": first_non_empty(d.get("radiologist_report_by"), "the diagnostic radiologist"),
        "P062_01": first_non_empty(d.get("educational_psychologist_by"), "the educational psychologist"),
        "P063_01": first_non_empty(d.get("occupational_therapist_by"), "the occupational therapist"),
        "P064_01": claimant,
        "P065_01": claimant,
        "P066_01": guardian_ref,
        "P070_01": claimant,
        "P070_02": (d.get("current_age", "").replace(" years", "") or ""),
        "P070_03": family_desc,
        "P075_01": claimant,
        "P075_02": first_non_empty(d.get("people_in_house"), guardian),
        "P075_03": "",
        "P078_01": guardian_ref,
        "P078_02": claimant,
        "P078_03": first_non_empty(d.get("accident_ordinal"), "first"),
        "P078_04": d.get("claimant_pronoun_object", "her"),
        "P078_05": first_non_empty(d.get("other_accidents_count"), "no"),
        "P078_06": guardian_ref,
        "P078_07": claimant,
        "P078_08": first_non_empty(d.get("medical_history_prior"), "no significant pre-existing medical history was reported"),
        "P078_09": first_non_empty(d.get("developmental_source"), "The intake information"),
        "P078_10": d.get("developmental_source_page", ""),
        "P078_11": d.get("developmental_source_report", ""),
        "P078_12": claimant,
        "P078_13": first_non_empty(d.get("developmental_milestones"), "were reportedly achieved within normal developmental timeframes"),
        "P081_01": claimant,
        "P082_01": d.get("education_nb_text", ""),
        "P085_01": claimant,
        "P085_02": first_non_empty(d.get("school_reports_pre_from"), "the available school records"),
        "P085_03": d.get("claimant_pronoun_possessive", "her"),
        "P085_04": first_non_empty(d.get("school_reports_post_from"), "the available post-accident school records"),
        "P085_05": d.get("claimant_pronoun_possessive", "her"),
        "P085_06": guardian_ref,
        "P085_07": guardian_ref,
        "P085_08": d.get("collateral_has_contacts", "not"),
        "P085_09": claimant,
        "P088_01": guardian_ref,
        "P088_02": format_date(d.get("date_accident", "")),
        "P088_03": claimant,
        "P088_04": first_non_empty(d.get("type_of_accident"), "Pedestrian").split()[0],
        "P088_05": guardian_ref,
        "P088_06": claimant,
        "P088_07": claimant,
        "P088_08": first_non_empty(d.get("first_hospital"), "hospital"),
        "P088_09": d.get("referred_hospital", ""),
        "P088_10": first_non_empty(d.get("hospital_records_from"), d.get("first_hospital"), ""),
        "P088_11": claimant,
        "P088_12": d.get("hospital_admission_date", ""),
        "P088_13": first_non_empty(d.get("hospital_accident_follow_up"), "received treatment"),
        "P091_01": claimant,
        "P092_01": ensure_sentence(d.get("injury_1", "")),
        "P095_01": d.get("treatment_1", ""),
        "P096_01": d.get("treatment_2", ""),
        "P097_01": d.get("treatment_3", ""),
        "P098_01": d.get("treatment_4", ""),
        "P099_01": d.get("treatment_5", ""),
        "P101_01": claimant,
        "P103_01": d.get("complaint_1", ""),
        "P104_01": d.get("complaint_2", ""),
        "P105_01": d.get("complaint_3", ""),
        "P106_01": d.get("complaint_4", ""),
        "P107_01": d.get("complaint_5", ""),
        "P108_01": d.get("complaint_6", ""),
        "P109_01": d.get("complaint_7", ""),
        "P111_01": ensure_sentence(first_non_empty(d.get("pre_morbid_potential"), "The pre-morbid potential should be interpreted with reference to the educational and family background captured during intake")),
        "P112_01": ensure_sentence(first_non_empty(d.get("post_morbid_potential"), "The post-morbid potential should be interpreted with reference to the post-accident complaints and expert findings")),
        "P113_01": ensure_sentence(first_non_empty(d.get("loss_of_earnings_text"), "The loss of earnings opinion remains subject to the final educational, medical and actuarial findings")),
        "P115_01": ensure_sentence(first_non_empty(d.get("recommendation_1"), "Provision should be made for the recommended interventions and treatment proposed by the experts")),
        "P116_01": ensure_sentence(first_non_empty(d.get("recommendation_2"), "Additionally, provision should be made for any past and future medical expenses incurred")),
    }

    family_rows = d.get("family_members") or []
    for i in range(1, 4):
        values = family_rows[i - 1] if i - 1 < len(family_rows) else ["", "", "", "", ""]
        for j in range(1, 6):
            marker_map[f"FA_{i:02d}_{j:02d}"] = values[j - 1] if j - 1 < len(values) else ""

    edu_rows = d.get("education_history_all") or []
    for i in range(1, 12):
        values = edu_rows[i - 1] if i - 1 < len(edu_rows) else ["", "", "", ""]
        for j in range(1, 5):
            marker_map[f"ED_{i:02d}_{j:02d}"] = values[j - 1] if j - 1 < len(values) else ""

    client_info = {
        "CI_01": d.get("claimant_full_name", ""),
        "CI_02": d.get("id_passport", ""),
        "CI_03": format_date(d.get("dob", "")),
        "CI_05": d.get("nationality", ""),
        "CI_06": d.get("current_age", ""),
        "CI_07": d.get("gender", ""),
        "CI_08": d.get("home_language", ""),
        "CI_09": d.get("highest_education", ""),
        "CI_10": format_date(d.get("date_accident", "")),
        "CI_11": d.get("age_at_accident", ""),
        "CI_12": d.get("grade_at_accident", ""),
        "CI_13": d.get("current_grade", ""),
        "CI_14": d.get("type_of_accident", ""),
        "CI_15": format_date(d.get("assessment_date", "")),
        "CI_16": format_date(d.get("report_date", "")),
        "CI_17": d.get("attorney_name", ""),
        "CI_18": d.get("attorney_reference", ""),
        "CI_19": d.get("our_reference", ""),
    }
    marker_map.update(client_info)
    return marker_map

def populate_ip_sections(doc: Document, data: dict):
    tables = doc.tables
    if len(tables) >= 3:
        fill_client_info_table(tables[0], data)
        family_rows = data.get("family_members") or []
        fill_dynamic_table(tables[1], family_rows or [["", "", "", "", ""]], start_row=1, headers=data.get("family_members_headers"))
        edu_rows = data.get("education_history_all") or []
        fill_dynamic_table(tables[2], edu_rows or [["", "", "", ""]], start_row=1, headers=data.get("education_history_all_headers"))


def generate_ip_report_docx(template_path: str | Path, output_path: str | Path, data: dict):
    template_path = Path(template_path)
    output_path = Path(output_path)
    prepared = derive_ip_defaults(data)
    marker_map = build_marker_map(prepared)
    doc = Document(str(template_path))
    populate_ip_sections(doc, prepared)
    for paragraph in iter_all_paragraphs(doc):
        replace_markers_in_paragraph(paragraph, marker_map)

    overwrite_paragraph_after(doc, '1.3. Behaviour and Presentation', build_behaviour_paragraph(prepared))
    overwrite_document_received_lines(doc, build_documents_received(prepared))
    overwrite_paragraph_after(doc, '2.1. Personal and Family Background', f"{prepared.get('claimant_short_name') or prepared.get('claimant_full_name') or 'The claimant'} is {prepared.get('current_age') or 'a minor'} old. Familial details are detailed in Table 1 below.")
    overwrite_paragraph_after(doc, '2.2. Residential Information', build_residential_paragraph(prepared))
    overwrite_paragraph_after(doc, '2.3 Medical History', build_medical_history_paragraph(prepared))
    overwrite_paragraph_after(doc, '2.5 Collateral Information', build_collateral_paragraph(prepared))
    overwrite_paragraph_after(doc, 'THE ACCIDENT', build_accident_paragraph(prepared))
    rec1, rec2 = build_recommendation_paragraphs(prepared)
    overwrite_paragraph_after(doc, 'RECOMMENDATIONS', rec1)
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == 'RECOMMENDATIONS':
            count = 0
            for j in range(idx + 1, len(paragraphs)):
                if paragraphs[j].text.strip():
                    count += 1
                    if count == 2:
                        set_paragraph_text(paragraphs[j], rec2)
                    elif count == 3:
                        set_paragraph_text(paragraphs[j], 'Additionally, provision should also be made for any past and future medical expenses incurred.')
                    if count >= 3:
                        break
            break

    clear_doc_highlighting(doc)
    doc.save(str(output_path))

